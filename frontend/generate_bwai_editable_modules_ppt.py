from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


BASE_DIR = Path(__file__).resolve().parent
LOGO_DIR = BASE_DIR / "logo-cropped"
OUTPUT_PATH = BASE_DIR / "build-with-ai-modules-editable-v2.pptx"

SLIDE_W = Inches(10)
SLIDE_H = Inches(17.714)

COLORS = {
    "bg": "F3F1EC",
    "surface": "FFFDF9",
    "ink": "202124",
    "muted": "666B70",
    "line": "202124",
    "blue": "4285F4",
    "red": "EA4335",
    "yellow": "F9AB00",
    "yellow_soft": "FFE7A5",
    "green": "34A853",
    "panel": "FBFAF7",
    "panel_line": "DBD8D2",
}

FONT_CN = "Microsoft YaHei"
FONT_EN = "Arial"
FONT_HEAVY = "Arial Black"
FONT_MONO = "Consolas"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def blank_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(COLORS["bg"])
    return slide


def set_fill(shape, color: str | None):
    if color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(color)
    else:
        shape.fill.background()


def set_line(shape, color: str | None, width=2.0):
    if color:
        shape.line.color.rgb = rgb(color)
        shape.line.width = Pt(width)
    else:
        shape.line.fill.background()


def add_shape(slide, shape_type, x, y, w, h, fill=None, line=None, line_width=2.0, radius_text=None):
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shape, fill)
    set_line(shape, line, line_width)
    if radius_text:
        shape.adjustments[0] = radius_text
    return shape


def add_text(
    slide,
    x,
    y,
    w,
    h,
    text,
    size,
    font=FONT_CN,
    color="202124",
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def add_picture_contain(slide, path: Path, x, y, w, h):
    if not path.exists():
        return None
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    return pic


def add_pill(slide, x, y, w, h, text, fill, line, text_color, size=16, font=FONT_EN, bold=True):
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line, line_width=2.0)
    add_text(slide, x, y + 0.06, w, h - 0.08, text, size=size, font=font, color=text_color, bold=bold, align=PP_ALIGN.CENTER)


def draw_brace(slide, x, y, h, left=True):
    bar_w = 0.56
    arm_w = 0.18
    arm_h = 0.64
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, bar_w, h, fill=COLORS["yellow_soft"], line=COLORS["line"], line_width=2.2)
    mid_y = y + h / 2 - arm_h / 2
    arm_x = x + bar_w - 0.02 if left else x - arm_w + 0.02
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, arm_x, mid_y, arm_w, arm_h, fill=COLORS["yellow_soft"], line=COLORS["line"], line_width=2.2)


def draw_lockup(slide, x, y, scale=1.0):
    left_x = x
    brace_h = 1.9 * scale
    draw_brace(slide, left_x, y + 0.18 * scale, brace_h, left=True)
    add_text(slide, x + 0.95 * scale, y, 3.0 * scale, 0.95 * scale, "Build", size=44 * scale, font=FONT_HEAVY, color=COLORS["ink"], bold=True)
    add_text(slide, x + 0.95 * scale, y + 0.9 * scale, 3.25 * scale, 0.95 * scale, "with AI", size=44 * scale, font=FONT_HEAVY, color=COLORS["ink"], bold=True)

    add_shape(slide, MSO_AUTO_SHAPE_TYPE.HEXAGON, x + 3.72 * scale, y + 0.2 * scale, 0.62 * scale, 0.62 * scale, fill=COLORS["blue"], line=COLORS["line"], line_width=2.0)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.DIAMOND, x + 3.88 * scale, y + 0.36 * scale, 0.30 * scale, 0.30 * scale, fill=COLORS["surface"], line=COLORS["line"], line_width=1.8)

    draw_brace(slide, x + 5.0 * scale, y + 0.18 * scale, brace_h, left=False)

    add_pill(
        slide,
        x + 1.2 * scale,
        y + 2.0 * scale,
        2.9 * scale,
        0.5 * scale,
        "Wenzhou Kean University",
        COLORS["surface"],
        COLORS["line"],
        COLORS["ink"],
        size=12 * scale,
        font=FONT_CN,
    )


def draw_info_band(slide, x, y, w, h):
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, fill=COLORS["surface"], line=COLORS["panel_line"], line_width=1.6)
    col_w = w / 3
    for idx in (1, 2):
        line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x + col_w * idx), Inches(y + 0.16), Inches(0.01), Inches(h - 0.32))
        set_fill(line, COLORS["panel_line"])
        set_line(line, None)

    blocks = [
        ("这次不只是讲", "讲座 + Workshop + 实操产出", "从理解、实操到产出，完整走一遍 AI 落地过程。"),
        ("Audience", "开放参与", "全校及社会所有对 AI 感兴趣的人，不限专业背景。"),
        ("Why build", "真正落地", "从科研工作流到模型微调与部署，让 AI 走到可运行的应用。"),
    ]

    for i, (label, title, desc) in enumerate(blocks):
        bx = x + col_w * i + 0.18
        add_text(slide, bx, y + 0.16, col_w - 0.28, 0.28, label, 10, font=FONT_EN, color=COLORS["muted"], bold=False)
        add_text(slide, bx, y + 0.34, col_w - 0.28, 0.44, title, 18, font=FONT_CN, color=COLORS["ink"], bold=True)
        add_text(slide, bx, y + 0.78, col_w - 0.28, 0.36, desc, 9.5, font=FONT_CN, color=COLORS["muted"])


def draw_card(slide, x, y, w, h, accent, label, title, subtitle, items):
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, fill=COLORS["surface"], line=COLORS["line"], line_width=2.2)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, 0.16, fill=accent, line=None)
    add_pill(slide, x + 0.18, y + 0.22, 0.95, 0.42, label, COLORS["surface"], COLORS["line"], accent, size=12, font=FONT_EN)
    add_text(slide, x + 0.18, y + 0.74, w - 0.36, 0.54, title, 26, font=FONT_CN, color=COLORS["ink"], bold=True)
    add_text(slide, x + 0.18, y + 1.22, w - 0.36, 0.24, subtitle, 10.5, font=FONT_MONO, color=COLORS["muted"])

    current_y = y + 1.64
    for heading, body in items:
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 0.2, current_y + 0.07, 0.12, 0.12, fill=accent, line=None)
        add_text(slide, x + 0.4, current_y, w - 0.6, 0.28, heading, 15, font=FONT_CN, color=COLORS["ink"], bold=True)
        add_text(slide, x + 0.4, current_y + 0.28, w - 0.62, 0.44, body, 10.5, font=FONT_CN, color=COLORS["muted"])
        current_y += 0.78


def draw_org_band(slide, x, y, w, h):
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, fill=COLORS["surface"], line=COLORS["panel_line"], line_width=1.8)
    add_text(slide, x + 0.18, y + 0.12, 1.1, 0.34, "联合发起", 18, font=FONT_CN, color=COLORS["ink"], bold=True)
    add_text(slide, x + 1.2, y + 0.16, 1.5, 0.22, "CO-ORGANIZED BY", 10, font=FONT_EN, color=COLORS["muted"], bold=False)

    logos = [
        (LOGO_DIR / "2.png", "温州肯恩理工学院\n计算机科学与技术系", False),
        (LOGO_DIR / "3.png", "Google Developer Groups\n温州站", False),
        (LOGO_DIR / "4.png", "温州市软件行业协会\n人工智能专业委员会", False),
        (LOGO_DIR / "datawhale-mark.png", "Datawhale\n开源学习组织", True),
        (LOGO_DIR / "5.png", "温州肯恩大学\nAI 社团", False),
    ]

    tile_w = (w - 0.36 - 0.14 * 4) / 5
    start_x = x + 0.18
    for i, (logo_path, caption, dark_bg) in enumerate(logos):
        tx = start_x + i * (tile_w + 0.14)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, tx, y + 0.5, tile_w, h - 0.64, fill=COLORS["surface"], line=COLORS["line"], line_width=1.8)
        if dark_bg:
            add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, tx + 0.12, y + 0.65, tile_w - 0.24, 0.78, fill="1E2D42", line=None)
            if logo_path.exists():
                add_picture_contain(slide, logo_path, tx + 0.32, y + 0.72, tile_w - 0.64, 0.64)
        else:
            if logo_path.exists():
                add_picture_contain(slide, logo_path, tx + 0.25, y + 0.66, tile_w - 0.5, 0.72)
        add_text(slide, tx + 0.08, y + 1.5, tile_w - 0.16, 0.42, caption, 9.4, font=FONT_CN, color=COLORS["muted"], bold=False, align=PP_ALIGN.CENTER)


def draw_module_header(slide, title, subtitle):
    add_text(slide, 0.56, 0.36, 4.8, 0.46, title, 24, font=FONT_CN, color=COLORS["ink"], bold=True)
    add_text(slide, 0.56, 0.82, 4.8, 0.26, subtitle, 11, font=FONT_MONO, color=COLORS["muted"], bold=False)


def build_full_slide(prs):
    slide = blank_slide(prs)
    draw_lockup(slide, 1.1, 0.25, scale=1.0)
    draw_info_band(slide, 0.32, 3.18, 9.36, 1.12)
    add_pill(slide, 4.2, 4.48, 1.6, 0.34, "Not just talk. Build on site.", COLORS["surface"], COLORS["line"], COLORS["ink"], size=9, font=FONT_MONO)

    draw_card(
        slide, 5.35, 5.0, 4.15, 3.55, COLORS["blue"], "Opening", "活动开场", "KICKOFF & OVERVIEW",
        [
            ("领导嘉宾发言", "邀请领导与嘉宾为本次 Build with AI 致辞，建立现场氛围并点明活动主题。"),
            ("AI Agent 引入", "通过贴近场景的方式引入 AI agent 概念，帮助大家快速进入今天的主题语境。"),
            ("活动内容概述", "介绍整场活动结构、互动方式与现场实践安排，让参与路径更清晰。"),
            ("AI 知识分享", "建立基础认知，帮助不同背景的参与者都能顺畅进入后续内容。"),
        ]
    )
    draw_card(
        slide, 0.45, 6.15, 4.25, 2.7, COLORS["yellow"], "Sharing", "AI 知识分享", "FRONTIER + RESEARCH WORKFLOW",
        [
            ("Harness Engineering 深度分享", "从工程视角拆解 AI 系统如何真正被构建出来，看懂能力、架构与落地流程。"),
            ("Claude Code 科研实战", "从文献综述到研究洞察，一整套工作流现场跑通，看到 AI 如何服务学习与研究。"),
        ]
    )
    draw_card(
        slide, 5.35, 8.95, 4.15, 2.7, COLORS["green"], "Experience", "茶歇与体验", "CONNECT & INTERACT",
        [
            ("开发者面对面交流", "和讲者、开发者、同学直接交流，获得真实经验与灵感。"),
            ("技术体验区 & AI 互动区", "边聊边试、真正参与，现场上手体验各类 AI 应用。"),
        ]
    )
    draw_card(
        slide, 0.45, 10.45, 4.45, 4.0, COLORS["red"], "Build", "项目实战", "HANDS-ON BUILDING",
        [
            ("Gemma4 × 多设备智能体", "手机 + 大模型协同，让 AI 从“工具”变成“系统”，看到更完整的交互方式。"),
            ("Fine Tuning Gemma", "让模型进一步适配你的数据、任务和具体场景。"),
            ("OpenClaw 部署实操", "从部署到运行，完成一个真正可以跑起来、可以展示的 AI 应用。"),
        ]
    )
    draw_org_band(slide, 0.28, 15.0, 9.44, 2.14)


def build_title_slide(prs):
    slide = blank_slide(prs)
    draw_module_header(slide, "标题模块", "Editable title lockup")
    draw_lockup(slide, 1.05, 2.0, scale=1.1)
    add_text(slide, 1.2, 5.45, 7.6, 0.5, "这一页的标题、学校名、徽章和左右括号都可以单独编辑或替换。", 14, font=FONT_CN, color=COLORS["muted"], align=PP_ALIGN.CENTER)


def build_info_slide(prs):
    slide = blank_slide(prs)
    draw_module_header(slide, "信息栏模块", "Editable intro information band")
    draw_info_band(slide, 0.5, 2.0, 9.0, 1.3)
    add_text(slide, 0.8, 4.0, 8.4, 0.5, "适合自己修改为时间、地点、面向对象、活动卖点等概述信息。", 14, font=FONT_CN, color=COLORS["muted"], align=PP_ALIGN.CENTER)


def build_opening_slide(prs):
    slide = blank_slide(prs)
    draw_module_header(slide, "开场模块", "Editable opening card")
    draw_card(
        slide, 1.0, 2.0, 8.0, 5.1, COLORS["blue"], "Opening", "活动开场", "KICKOFF & OVERVIEW",
        [
            ("领导嘉宾发言", "邀请领导与嘉宾为本次 Build with AI 致辞，建立现场氛围并点明活动主题。"),
            ("AI Agent 引入", "通过贴近场景的方式引入 AI agent 概念，帮助大家快速进入今天的主题语境。"),
            ("活动内容概述", "介绍整场活动结构、互动方式与现场实践安排，让参与路径更清晰。"),
            ("AI 知识分享", "建立基础认知，帮助不同背景的参与者都能顺畅进入后续内容。"),
        ]
    )


def build_share_slide(prs):
    slide = blank_slide(prs)
    draw_module_header(slide, "分享模块", "Editable AI knowledge sharing card")
    draw_card(
        slide, 1.0, 2.1, 8.0, 3.2, COLORS["yellow"], "Sharing", "AI 知识分享", "FRONTIER + RESEARCH WORKFLOW",
        [
            ("Harness Engineering 深度分享", "从工程视角拆解 AI 系统如何真正被构建出来，看懂能力、架构与落地流程。"),
            ("Claude Code 科研实战", "从文献综述到研究洞察，一整套工作流现场跑通，看到 AI 如何服务学习与研究。"),
        ]
    )


def build_experience_slide(prs):
    slide = blank_slide(prs)
    draw_module_header(slide, "体验模块", "Editable experience card")
    draw_card(
        slide, 1.0, 2.1, 8.0, 3.1, COLORS["green"], "Experience", "茶歇与体验", "CONNECT & INTERACT",
        [
            ("开发者面对面交流", "和讲者、开发者、同学直接交流，获得真实经验与灵感。"),
            ("技术体验区 & AI 互动区", "边聊边试、真正参与，现场上手体验各类 AI 应用。"),
        ]
    )


def build_build_slide(prs):
    slide = blank_slide(prs)
    draw_module_header(slide, "实战模块", "Editable hands-on building card")
    draw_card(
        slide, 0.9, 1.9, 8.2, 5.05, COLORS["red"], "Build", "项目实战", "HANDS-ON BUILDING",
        [
            ("Gemma4 × 多设备智能体", "手机 + 大模型协同，让 AI 从“工具”变成“系统”，看到更完整的交互方式。"),
            ("Fine Tuning Gemma", "让模型进一步适配你的数据、任务和具体场景。"),
            ("OpenClaw 部署实操", "从部署到运行，完成一个真正可以跑起来、可以展示的 AI 应用。"),
        ]
    )


def build_org_slide(prs):
    slide = blank_slide(prs)
    draw_module_header(slide, "联合发起模块", "Editable organizer logo strip")
    draw_org_band(slide, 0.35, 2.0, 9.3, 2.3)


def build_assets_slide(prs):
    slide = blank_slide(prs)
    draw_module_header(slide, "Logo Assets", "Editable logo source page")
    logos = [
        ("WKU · CSMT", LOGO_DIR / "2.png"),
        ("GDG Wenzhou", LOGO_DIR / "3.png"),
        ("WZSIA · AI", LOGO_DIR / "4.png"),
        ("Datawhale", LOGO_DIR / "datawhale-mark.png"),
        ("AI Club", LOGO_DIR / "5.png"),
    ]
    start_x = 0.55
    tile_w = 1.75
    for i, (label, path) in enumerate(logos):
        x = start_x + i * 1.88
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, 2.1, tile_w, 2.0, fill=COLORS["surface"], line=COLORS["line"], line_width=1.8)
        if label == "Datawhale":
            add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x + 0.12, 2.25, tile_w - 0.24, 0.92, fill="1E2D42", line=None)
            add_picture_contain(slide, path, x + 0.35, 2.32, tile_w - 0.7, 0.78)
        else:
            add_picture_contain(slide, path, x + 0.16, 2.26, tile_w - 0.32, 0.82)
        add_text(slide, x + 0.08, 3.28, tile_w - 0.16, 0.32, label, 10.5, font=FONT_EN, color=COLORS["ink"], align=PP_ALIGN.CENTER, bold=True)


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    build_full_slide(prs)
    build_title_slide(prs)
    build_info_slide(prs)
    build_opening_slide(prs)
    build_share_slide(prs)
    build_experience_slide(prs)
    build_build_slide(prs)
    build_org_slide(prs)
    build_assets_slide(prs)

    prs.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
