from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


BASE_DIR = Path(__file__).resolve().parent
ASSETS_DIR = BASE_DIR / "assets"
LOGO_DIR = BASE_DIR / "logo-cropped"
OUTPUT_PATH = BASE_DIR / "build-with-ai-poster-editable-v1.pptx"

SLIDE_W = 16
SLIDE_H = 10

COLORS = {
    "paper": "F5F3ED",
    "hero": "FCFBF8",
    "hero_line": "E7E8E4",
    "ink": "1F2023",
    "line": "2C2E33",
    "blue": "4F7FE0",
    "blue_deep": "174780",
    "blue_soft": "CAE8FB",
    "blue_panel": "E8F2FE",
    "blue_outline": "355F93",
    "sky": "72B5EA",
    "orange": "FFB20D",
    "orange_soft": "FFE9A0",
    "gray": "8A8A8A",
    "panel_line": "D8E3F0",
}


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.replace("#", ""))


def set_no_fill(shape):
    shape.fill.background()


def set_no_line(shape):
    shape.line.fill.background()


def add_shape(slide, shape_type, x, y, w, h, fill=None, line=None, line_width=1.5):
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    else:
        set_no_fill(shape)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        set_no_line(shape)
    return shape


def add_textbox(
    slide,
    x,
    y,
    w,
    h,
    text,
    font_size,
    font_name="Arial",
    color="1F2023",
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_no_fill(box)
    set_no_line(box)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = valign
    paragraphs = str(text).split("\n")
    for i, part in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = part
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def add_picture_fit(slide, path: Path, x, y, w, h):
    with Image.open(path) as img:
        img_w, img_h = img.size
    scale = min(w / img_w, h / img_h)
    final_w = img_w * scale
    final_h = img_h * scale
    left = x + (w - final_w) / 2
    top = y + (h - final_h) / 2
    return slide.shapes.add_picture(
        str(path), Inches(left), Inches(top), width=Inches(final_w), height=Inches(final_h)
    )


def add_fill_line(slide, x, y, w, h, color):
    return add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, w, h, fill=color, line=None)


def draw_spark_badge(slide, x, y, size):
    add_shape(
        slide,
        MSO_AUTO_SHAPE_TYPE.HEXAGON,
        x,
        y,
        size,
        size,
        fill=COLORS["blue"],
        line=COLORS["line"],
        line_width=2.2,
    )
    add_textbox(
        slide,
        x + size * 0.22,
        y + size * 0.16,
        size * 0.56,
        size * 0.56,
        "✦",
        font_size=28 * size,
        font_name="Arial",
        color="FFFDF6",
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def draw_grid_chip(slide, x, y, w, h):
    add_shape(
        slide,
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        x,
        y,
        w,
        h,
        fill=COLORS["blue_soft"],
        line=COLORS["line"],
        line_width=2.0,
    )
    dot_w = w * 0.19
    gap_x = w * 0.08
    gap_y = h * 0.14
    start_x = x + w * 0.15
    start_y = y + h * 0.15
    for row in range(2):
        for col in range(3):
            add_shape(
                slide,
                MSO_AUTO_SHAPE_TYPE.OVAL,
                start_x + col * (dot_w + gap_x),
                start_y + row * (dot_w + gap_y),
                dot_w,
                dot_w,
                fill=COLORS["blue"],
                line=COLORS["line"],
                line_width=1.7,
            )


def draw_title_module(slide, x, y, scale=1.0):
    add_shape(
        slide,
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        x + 0.18 * scale,
        y,
        3.88 * scale,
        0.56 * scale,
        fill=COLORS["orange"],
        line=COLORS["line"],
        line_width=2.0,
    )
    draw_grid_chip(slide, x + 3.26 * scale, y - 0.14 * scale, 1.78 * scale, 1.02 * scale)

    add_shape(
        slide,
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        x + 0.34 * scale,
        y + 0.44 * scale,
        4.62 * scale,
        3.76 * scale,
        fill="FFFFFF",
        line=COLORS["line"],
        line_width=2.2,
    )

    add_textbox(
        slide,
        x + 0.74 * scale,
        y + 1.16 * scale,
        2.52 * scale,
        1.84 * scale,
        "Build\nwith AI",
        font_size=52 * scale,
        font_name="Arial Black",
        color=COLORS["ink"],
        bold=True,
    )
    draw_spark_badge(slide, x + 3.5 * scale, y + 1.28 * scale, 0.9 * scale)

    add_shape(
        slide,
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        x + 1.28 * scale,
        y + 3.22 * scale,
        1.44 * scale,
        0.84 * scale,
        fill=COLORS["orange_soft"],
        line=COLORS["line"],
        line_width=2.0,
    )
    add_picture_fit(
        slide,
        LOGO_DIR / "3.png",
        x + 1.4 * scale,
        y + 3.28 * scale,
        1.18 * scale,
        0.64 * scale,
    )

    add_shape(
        slide,
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        x + 2.72 * scale,
        y + 3.22 * scale,
        1.52 * scale,
        0.84 * scale,
        fill="FFFFFF",
        line=COLORS["line"],
        line_width=2.0,
    )
    add_textbox(
        slide,
        x + 2.95 * scale,
        y + 3.42 * scale,
        1.06 * scale,
        0.42 * scale,
        "2026",
        font_size=28 * scale,
        font_name="Arial",
        color=COLORS["ink"],
        bold=False,
        align=PP_ALIGN.CENTER,
    )

    add_shape(
        slide,
        MSO_AUTO_SHAPE_TYPE.OVAL,
        x + 4.1 * scale,
        y + 3.3 * scale,
        0.64 * scale,
        0.64 * scale,
        fill=COLORS["sky"],
        line=COLORS["line"],
        line_width=2.0,
    )
    add_shape(
        slide,
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        x + 2.64 * scale,
        y + 4.5 * scale,
        2.28 * scale,
        0.38 * scale,
        fill=COLORS["orange"],
        line=COLORS["line"],
        line_width=1.8,
    )


def draw_campus_lockup(slide, x, y, scale=1.0):
    add_picture_fit(slide, ASSETS_DIR / "kean-seal-crop.png", x, y, 0.88 * scale, 0.88 * scale)
    add_textbox(
        slide,
        x + 0.96 * scale,
        y + 0.02 * scale,
        3.0 * scale,
        0.34 * scale,
        "温州肯恩大学",
        font_size=27 * scale,
        font_name="SimSun",
        color=COLORS["blue_deep"],
        bold=True,
    )
    add_textbox(
        slide,
        x + 0.98 * scale,
        y + 0.38 * scale,
        3.0 * scale,
        0.24 * scale,
        "WENZHOU-KEAN UNIVERSITY",
        font_size=12 * scale,
        font_name="Times New Roman",
        color=COLORS["blue_deep"],
        bold=False,
    )


def draw_info_block(slide, x, y, scale=1.0):
    add_textbox(
        slide,
        x,
        y,
        4.8 * scale,
        0.28 * scale,
        "BUILD WITH AI - WENZHOU EDITION",
        font_size=15 * scale,
        font_name="Arial",
        color="2B5A95",
        bold=True,
    )
    add_fill_line(slide, x, y + 0.42 * scale, 4.46 * scale, 0.03 * scale, COLORS["line"])

    add_textbox(
        slide,
        x,
        y + 0.72 * scale,
        2.0 * scale,
        0.24 * scale,
        "LOCATION",
        font_size=16 * scale,
        font_name="Arial",
        color=COLORS["gray"],
        bold=True,
    )
    add_textbox(
        slide,
        x,
        y + 1.04 * scale,
        4.5 * scale,
        1.0 * scale,
        "Wenzhou-Kean\nUniversity",
        font_size=36 * scale,
        font_name="Arial Black",
        color=COLORS["ink"],
        bold=True,
    )

    add_textbox(
        slide,
        x,
        y + 2.18 * scale,
        1.4 * scale,
        0.24 * scale,
        "DATE",
        font_size=16 * scale,
        font_name="Arial",
        color=COLORS["gray"],
        bold=True,
    )
    add_textbox(
        slide,
        x,
        y + 2.5 * scale,
        4.5 * scale,
        0.52 * scale,
        "April 26, 2026",
        font_size=30 * scale,
        font_name="Arial Black",
        color=COLORS["ink"],
        bold=True,
    )
    add_fill_line(slide, x, y + 3.18 * scale, 4.46 * scale, 0.03 * scale, COLORS["line"])


def draw_logo_tile(slide, x, y, w, h, image_path=None, text=None):
    add_shape(
        slide,
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        x,
        y,
        w,
        h,
        fill="FFFFFF",
        line=COLORS["panel_line"],
        line_width=1.0,
    )
    if image_path:
        add_picture_fit(slide, image_path, x + 0.08, y + 0.04, w - 0.16, h - 0.08)
    elif text:
        add_textbox(
            slide,
            x + 0.08,
            y + 0.16,
            w - 0.16,
            h - 0.2,
            text,
            font_size=14,
            font_name="Arial",
            color=COLORS["blue_deep"],
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )


def draw_partner_section(slide, x, y, w, h, scale=1.0):
    add_shape(
        slide,
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        x,
        y,
        w,
        h,
        fill="FFFFFF",
        line=COLORS["panel_line"],
        line_width=1.2,
    )

    add_textbox(
        slide,
        x + 0.18 * scale,
        y + 0.2 * scale,
        1.0 * scale,
        0.34 * scale,
        "活动单位",
        font_size=18 * scale,
        font_name="Microsoft YaHei",
        color="5F7EA7",
        bold=True,
    )
    add_textbox(
        slide,
        x + 0.18 * scale,
        y + 0.5 * scale,
        1.0 * scale,
        0.16 * scale,
        "EVENT PARTNERS",
        font_size=8 * scale,
        font_name="Arial",
        color="5F7EA7",
        bold=True,
    )

    tile_y = y + 0.14 * scale
    tile_w = 1.75 * scale
    tile_h = 0.44 * scale
    start_x = x + 1.72 * scale
    gap = 0.16 * scale

    draw_logo_tile(slide, start_x + 0 * (tile_w + gap), tile_y, tile_w, tile_h, LOGO_DIR / "3.png")
    draw_logo_tile(
        slide,
        start_x + 1 * (tile_w + gap),
        tile_y,
        tile_w,
        tile_h,
        LOGO_DIR / "datawhale-mark-trimmed.png",
    )
    draw_logo_tile(slide, start_x + 2 * (tile_w + gap), tile_y, tile_w, tile_h, LOGO_DIR / "4.png")
    draw_logo_tile(slide, start_x + 3 * (tile_w + gap), tile_y, tile_w, tile_h, LOGO_DIR / "2.png")
    draw_logo_tile(slide, start_x + 4 * (tile_w + gap), tile_y, tile_w, tile_h, LOGO_DIR / "5.png")

    add_fill_line(slide, x + 0.16 * scale, y + 0.74 * scale, w - 0.32 * scale, 0.015 * scale, "E5EAF1")

    label_x = x + 1.6 * scale
    content_x = x + 2.32 * scale
    row_y = y + 0.84 * scale
    line_h = 0.16 * scale
    body_size = 8.4 * scale

    labels = ["主办单位", "承办单位", "协办单位"]
    values = [
        "Google Developer Groups Wenzhou（谷歌开发者社区温州站）、Datawhale",
        "温州市软件行业协会人工智能专业委员会、温州肯恩理工学院计算机科学与技术系",
        "温州肯恩大学 AI 社团",
    ]
    for idx, (label, value) in enumerate(zip(labels, values)):
        line_top = row_y + idx * 0.18 * scale
        add_textbox(
            slide,
            label_x,
            line_top,
            0.62 * scale,
            line_h,
            label,
            font_size=8.4 * scale,
            font_name="Microsoft YaHei",
            color="1F528F",
            bold=True,
        )
        add_textbox(
            slide,
            content_x,
            line_top,
            w - (content_x - x) - 0.2 * scale,
            line_h,
            value,
            font_size=body_size,
            font_name="Microsoft YaHei",
            color="244260",
            bold=False,
        )


def draw_full_poster_slide(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(COLORS["paper"])

    add_shape(
        slide,
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        0.54,
        0.56,
        14.92,
        7.42,
        fill=COLORS["hero"],
        line=COLORS["hero_line"],
        line_width=1.2,
    )
    add_shape(
        slide,
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        0.82,
        0.82,
        14.36,
        6.9,
        fill=None,
        line="EEF0EB",
        line_width=0.8,
    )

    draw_title_module(slide, 0.72, 1.14, 1.0)
    slide.shapes.add_picture(str(ASSETS_DIR / "campus-illustration.png"), Inches(9.9), Inches(4.1), width=Inches(4.8))
    draw_campus_lockup(slide, 10.8, 0.68, 1.0)
    draw_info_block(slide, 8.7, 1.82, 1.0)
    draw_partner_section(slide, 0.54, 8.16, 14.92, 1.58, 1.0)


def draw_component_slide(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("FBFBF9")

    add_textbox(
        slide,
        0.66,
        0.42,
        5.5,
        0.4,
        "Editable Build With AI Poster Blocks",
        font_size=26,
        font_name="Arial Black",
        color=COLORS["ink"],
        bold=True,
    )
    add_textbox(
        slide,
        0.66,
        0.82,
        6.8,
        0.22,
        "Slide 1 is the full poster. This slide keeps the major blocks separated for easier editing and rearranging.",
        font_size=10,
        font_name="Arial",
        color="6E7782",
        bold=False,
    )

    draw_campus_lockup(slide, 0.72, 1.22, 0.95)
    draw_info_block(slide, 8.64, 1.34, 0.82)
    draw_title_module(slide, 0.54, 2.12, 0.78)

    add_textbox(
        slide,
        8.6,
        5.0,
        2.2,
        0.26,
        "Campus Illustration",
        font_size=14,
        font_name="Arial",
        color="5F7EA7",
        bold=True,
    )
    slide.shapes.add_picture(
        str(ASSETS_DIR / "campus-illustration.png"),
        Inches(9.0),
        Inches(5.28),
        width=Inches(5.3),
    )

    draw_partner_section(slide, 0.54, 7.68, 14.92, 1.58, 1.0)


def main():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    draw_full_poster_slide(slide1)

    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    draw_component_slide(slide2)

    prs.save(OUTPUT_PATH)
    print(f"saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
